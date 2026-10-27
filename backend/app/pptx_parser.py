"""Bounded PPTX normalization; never extract ZIP members or load external relationships."""
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
from fastapi import HTTPException
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml import etree

MAX_PPTX = 25 * 1024 * 1024


def parse_pptx(content, snapshot_id):
    try:
        if len(content) > MAX_PPTX: raise ValueError("PPTX превышает 25 MiB")
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                raise ValueError("Распакованный PPTX превышает лимиты (100 MiB / 2000 файлов)")
            if len({e.filename for e in entries}) != len(entries): raise ValueError("Повторяющиеся файлы ZIP")
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in entry.filename or entry.flag_bits & 1:
                    raise ValueError("Недопустимая структура ZIP")
                if entry.file_size > max(1, entry.compress_size) * 1000: raise ValueError("Слишком высокая степень сжатия ZIP")
                if entry.filename.endswith((".xml", ".rels")):
                    xml = archive.read(entry)
                    root = etree.fromstring(xml, parser=etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False))
                    if root.getroottree().docinfo.doctype: raise ValueError("DTD и XML entities запрещены")
            if "ppt/presentation.xml" not in archive.namelist(): raise ValueError("Файл не является PPTX")
        presentation = Presentation(BytesIO(content))
        if len(presentation.slides) > 300: raise ValueError("Максимум 300 слайдов")
        blocks, tables, warnings = [], [], []
        characters = 0

        def block(text, slide, shape, *, cell=None, heading=False):
            nonlocal characters
            text = text.strip()
            if not text: return None
            characters += len(text)
            if len(blocks) >= 20000 or characters > 2_000_000: raise ValueError("Слишком много текста в PPTX")
            item = {"id": f"b{len(blocks)}", "kind": "heading" if heading else "text", "level": 1 if heading else 0,
                    "text": text, "slide_index": slide, "shape_id": shape}
            if cell is not None: item["cell"] = cell
            blocks.append(item)
            return item

        def walk(shapes, slide, depth=0):
            if depth > 20: raise ValueError("Слишком глубокая группа фигур")
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes, slide, depth + 1)
                elif shape.has_table:
                    rows = []
                    for ri, row in enumerate(shape.table.rows):
                        if any(c.is_spanned or c.is_merge_origin for c in row.cells):
                            warnings.append(f"Слайд {slide}: строка {ri + 1} с объединёнными ячейками пропущена")
                            continue
                        rows.append([block(c.text, slide, shape.shape_id, cell=[ri, ci]) for ci, c in enumerate(row.cells)])
                    tables.append({"id": f"t{len(tables)}", "rows": rows, "slide_index": slide, "shape_id": shape.shape_id})
                elif shape.has_text_frame:
                    is_title = shape == presentation.slides[slide - 1].shapes.title
                    for paragraph in shape.text_frame.paragraphs:
                        block(paragraph.text, slide, shape.shape_id, heading=is_title)
                else:
                    warnings.append(f"Слайд {slide}, фигура {shape.shape_id}: изображение, диаграмма или SmartArt не распознаны")

        for index, slide in enumerate(presentation.slides, 1):
            count = len(blocks)
            walk(slide.shapes, index)
            if len(blocks) == count: warnings.append(f"Слайд {index}: нет доступного текста; OCR не выполняется")
        return {"contract_version": 1, "format": "pptx", "snapshot_id": snapshot_id, "blocks": blocks,
                "tables": tables, "warnings": warnings, "slide_count": len(presentation.slides)}
    except HTTPException: raise
    except Exception as exc:
        detail = str(exc) if isinstance(exc, ValueError) else "Повреждённый или неподдерживаемый PPTX"
        raise HTTPException(422, detail) from None
