from io import BytesIO
from zipfile import ZipFile
import pytest
import httpx
from pptx import Presentation
from pptx.util import Inches
from fastapi import HTTPException
from backend.app.pptx_parser import parse_pptx
from backend.app.confluence.attachments import resolve_attachment, fetch_attachment
from backend.app.confluence.client import ConfluenceClient, ConfluenceError


def presentation_bytes():
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    title=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(.5))
    title.text="Проблематика"
    body=slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1))
    body.text="Потери сообщений"
    end=slide.shapes.add_textbox(Inches(1), Inches(3), Inches(7), Inches(.5))
    end.text="Задачи проекта"
    group=slide.shapes.add_group_shape()
    group.shapes.add_textbox(0,0,Inches(2),Inches(1)).text="Вложенный текст"
    table=slide.shapes.add_table(2,2,Inches(1),Inches(4),Inches(7),Inches(2)).table
    table.cell(0,0).text="Название"; table.cell(0,1).text="Доставка"
    table.cell(1,0).text="Владелец"; table.cell(1,1).text="Команда А"
    prs.slides.add_slide(prs.slide_layouts[6])
    output=BytesIO();prs.save(output);return output.getvalue()


def test_parse_text_tables_groups_and_empty_slide():
    doc=parse_pptx(presentation_bytes(),"s")
    assert any(b["text"]=="Вложенный текст" for b in doc["blocks"])
    assert doc["tables"][0]["rows"][0][1]["text"]=="Доставка"
    assert all(b["slide_index"]==1 and b["shape_id"] for b in doc["blocks"])
    assert any("2" in w and "текст" in w for w in doc["warnings"])


def test_zip_validation():
    with pytest.raises(HTTPException):parse_pptx(b"broken", "s")
    buf=BytesIO()
    with ZipFile(buf,"w") as z:z.writestr("../escape","x")
    with pytest.raises(HTTPException):parse_pptx(buf.getvalue(),"s")


def test_resolve_and_download_attachment_via_api():
    seen=[];content=presentation_bytes()
    def handle(req):
        seen.append(req)
        assert req.headers["authorization"]=="Bearer test-pat"
        if req.url.path.endswith("/child/attachment"):
            assert req.url.params["filename"]=="Архитектура.pptx"
            return httpx.Response(200,json={"results":[{"id":"att77","title":"Архитектура.pptx","type":"attachment","version":{"number":3},"container":{"id":"42"},"_links":{"download":"/download/attachments/42/Architecture.pptx?version=3"}}]})
        if "/download/attachments/" in req.url.path:return httpx.Response(200,content=content)
        raise AssertionError(req.url)
    with ConfluenceClient("https://wiki.example/confluence","test-pat",transport=httpx.MockTransport(handle)) as client:
        meta=resolve_attachment(client,"https://wiki.example/confluence/download/attachments/42/%D0%90%D1%80%D1%85%D0%B8%D1%82%D0%B5%D0%BA%D1%82%D1%83%D1%80%D0%B0.pptx")
        assert meta["attachment_id"]=="att77" and meta["parent_page_id"]=="42" and meta["version"]==3
        assert fetch_attachment(client,meta)==content
    assert seen[-1].url.path.startswith("/confluence/download/attachments/42/")


def test_attachment_rejects_untrusted_link_and_redirect():
    with ConfluenceClient("https://wiki.example","x",transport=httpx.MockTransport(lambda r:httpx.Response(302,headers={"location":"https://evil.example/a"}))) as client:
        with pytest.raises(ConfluenceError):resolve_attachment(client,"https://evil.example/download/attachments/1/a.pptx")
        with pytest.raises(ConfluenceError):fetch_attachment(client,{"download_url":"https://wiki.example/download/attachments/1/a.pptx"})


def test_local_pptx_upload_review_evidence(client,type_id):
    r=client.post("/api/pptx/uploads",data={"type_id":type_id,"mode":"table"},files={"file":("decision.pptx",presentation_bytes(),"application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert r.status_code==201,r.text
    c=r.json(); assert c["proposals"]["title"]["value"]=="Доставка"
    ev=c["proposals"]["title"]["evidence"][0]
    assert ev["slide_index"]==1 and ev["shape_id"]
    assert client.get("/api/artifacts").json()==[]


def test_remote_attachment_supplements_parent_and_repeat_is_idempotent(tmp_path):
    from tests.test_confluence_jobs import setup_client, config, handler, enqueue
    content=presentation_bytes()
    def remote(req):
        if req.url.path.endswith("/child/attachment"):
            return httpx.Response(200,json={"results":[{"id":"att77","title":"Architecture.pptx","type":"attachment","version":{"number":3},"container":{"id":"1"},"_links":{"download":"/download/attachments/1/Architecture.pptx?version=3"}}]})
        if req.url.path.startswith("/download/attachments/"): return httpx.Response(200,content=content)
        return handler(req)
    app,client,_=setup_client(tmp_path,remote)
    with client:
        connection_id,type_id=config(client)
        body={"connection_id":connection_id,"type_id":type_id,"link":"https://wiki.example/download/attachments/1/Architecture.pptx","expected_version":3}
        assert client.post("/api/pptx/attachments",json=body).status_code==409
        job=enqueue(client,connection_id,type_id).json();app.state.confluence_worker.run_once()
        page_id=client.get('/api/confluence/jobs/'+job['id']).json()['items'][0]['candidate_id']
        artifact=client.post(f'/api/candidates/{page_id}/review',json={"expected_revision":0,"fields":{"title":"Page","owner":"Manual"},"status":"published"}).json()
        response=client.post('/api/pptx/attachments',json=body)
        assert response.status_code==201,response.text
        candidate=response.json()
        assert candidate['target_id']==artifact['id'] and candidate['current_fields']['owner']=='Manual'
        assert len(client.get('/api/artifacts').json())==1
        again=client.post('/api/pptx/attachments',json=body).json()
        assert again['id']==candidate['id']
        accepted=client.post(f"/api/candidates/{candidate['id']}/review",json={"expected_revision":1,"fields":{"title":"Page","owner":"Команда А"},"status":"published"}).json()
        assert len(accepted['sources'])==2
        assert {s['role'] for s in accepted['sources']}=={'primary','supplementary'}
        assert accepted['provenance']['owner']['evidence'][0]['slide_index']==1
        assert client.post('/api/pptx/attachments',json={**body,'expected_version':2}).status_code==409


def test_pptx_dtd_utf16_is_rejected():
    original=ZipFile(BytesIO(presentation_bytes()))
    buf=BytesIO()
    with ZipFile(buf,'w') as target:
        for name in original.namelist():
            data=original.read(name)
            if name=='docProps/core.xml':
                data='<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY custom "test">]><x>&custom;</x>'.encode('utf-16')
            target.writestr(name,data)
    with pytest.raises(HTTPException,match='DTD'):parse_pptx(buf.getvalue(),'s')
